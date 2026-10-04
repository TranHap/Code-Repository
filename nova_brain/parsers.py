"""Découpe chaque type de fichier en unités de preuve citables, chacune avec un repère.

Chaque parseur retourne (meta, unites) :
- meta   : {"doc_date", "title", "message_id"?, "attachments"?, "incomplet"?}
- unites : liste de {"suffix", "unit", "repere", "date_fait", "auteur", "texte"}

Un format inconnu lève NonSupporte : on n'invente pas de texte à partir d'octets binaires.
"""
import csv
import datetime as dt
import email
import html
import io
import re
import zipfile
from email import policy
from html.parser import HTMLParser

import fitz  # PyMuPDF
import openpyxl

from .dates import _MOIS_RE, date_en_tete, email_date, entete_de_jour, find_date

PLACEHOLDER = "[capture non transcrite]"
# À incrémenter quand le découpage ou les dates changent : les fichiers déjà ingérés seront re-parsés.
PARSER_VERSION = 4
TEXT_EXT = {"txt", "md", "log", "text", "json", "xml", "yaml", "yml", "ics", "vtt", "srt"}
IMAGE_EXT = {"png", "jpg", "jpeg", "gif", "webp", "bmp"}

# "15:22 Élodie : ..." (transcript) ou "09:12 - Alex : ..." (Teams) ou "15:23 [silence]".
# (?!:\d) exclut les horodatages HH:MM:SS des journaux techniques.
RE_TOUR = re.compile(r"^(\d{1,2}:\d{2})(?!:\d)\s*(?:-\s*)?(?:(?P<who>[^:\[\]]{1,40}?)\s*:\s+)?(?P<txt>.+)$")
# "17 sept 14:23 - Boréal : ..." ou "12 août - Mélissa : ..." (commentaire de ticket, heure facultative)
RE_COMM = re.compile(rf"^(?P<stamp>\d{{1,2}}\s+(?:{_MOIS_RE})\.?(?:\s+(?P<h>\d{{1,2}}:\d{{2}}))?)\s*-\s*(?P<who>[^:]{{1,40}}?)\s*:\s+(?P<txt>.+)$", re.I)


class NonSupporte(Exception):
    pass


def decode_text(raw: bytes) -> str:
    """UTF-8 (avec ou sans BOM), sinon Windows-1252 (exports Bloc-notes / Excel FR)."""
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def looks_binary(raw: bytes) -> bool:
    head = raw[:4096]
    if b"\x00" in head:
        return True
    try:
        head.decode("utf-8")
        return False
    except UnicodeDecodeError as e:
        return e.start < len(head) - 4  # coupure d'un caractère multi-octets en fin de bloc : texte


# ───────────────────────── Texte ─────────────────────────

def _doc_date(lines, n=6):
    for l in lines[:n]:
        if d := find_date(l):
            return d
    return None


def parse_lines(text: str, label="ligne"):
    """Fichiers texte (CR, transcripts, tickets, Teams, notes, Markdown, logs) : une unité par ligne.

    - Une ligne qui n'est qu'une date ("16 septembre 2026") change le jour des tours suivants (Teams multi-jours).
    - Une ligne sans horodatage qui suit immédiatement un tour de parole le prolonge (même orateur) ;
      une ligne vide termine le tour.
    """
    lines = text.splitlines()
    doc_date = _doc_date(lines)
    year = int(doc_date[:4]) if doc_date else None
    day = doc_date[:10] if doc_date else None
    title = next((l.strip("# *").strip() for l in lines if l.strip()), "")
    unites, prev_tour = [], None
    for i, raw in enumerate(lines, start=1):
        line = raw.strip()
        if not line:
            prev_tour = None
            continue
        clean = line.replace("**", "").lstrip("# ").strip()
        if jour := entete_de_jour(clean, year):
            day = jour   # nouveau jour dans un fil (Teams multi-jours)
        # Date du fait = date qui ouvre la ligne, sinon le jour courant du document.
        u = {"suffix": f"L{i}", "unit": label, "repere": f"{label} {i}",
             "date_fait": date_en_tete(clean, year) or day or doc_date, "auteur": None, "texte": clean}
        if m := RE_COMM.match(clean):
            d = find_date(m["stamp"], year)
            if d and doc_date and d < doc_date[:10]:   # ticket ouvert en décembre, commenté en janvier
                d = find_date(m["stamp"], year + 1)
            u.update(unit="commentaire", repere=f"comm. {m['stamp']} ({label} {i})",
                     date_fait=(f"{d}T{m['h']}" if m["h"] else d) if d else None, auteur=m["who"].strip(), texte=m["txt"])
            prev_tour = None
        elif (m := RE_TOUR.match(clean)) and day:
            u.update(unit="tour", repere=f"{m.group(1)} ({label} {i})", date_fait=f"{day}T{m.group(1)}",
                     auteur=(m["who"] or "").strip() or None, texte=m["txt"])
            u["_t"], u["_l0"] = m.group(1), i
            prev_tour = u
        elif prev_tour is not None:
            prev_tour["texte"] += "\n" + clean
            prev_tour["repere"] = f"{prev_tour['_t']} ({label}s {prev_tour['_l0']}-{i})"
            continue
        unites.append(u)
    for u in unites:
        u.pop("_t", None), u.pop("_l0", None)
    return {"doc_date": doc_date, "title": title}, unites


# ───────────────────────── Courriel ─────────────────────────

class _HTMLText(HTMLParser):
    BLOCK = {"p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "table"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out, self.skip = [], False

    def handle_starttag(self, tag, attrs):
        self.skip = tag in ("style", "script") or self.skip
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("style", "script"):
            self.skip = False

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def html_to_text(s: str) -> str:
    p = _HTMLText()
    p.feed(s)
    text = html.unescape("".join(p.out))
    return "\n".join(re.sub(r"[ \t]+", " ", l).strip() for l in text.splitlines() if l.strip())


def norm_message_id(mid) -> str | None:
    mid = (str(mid) if mid else "").strip().strip("<>").strip().lower()
    return mid or None


def _attachments(msg):
    """Pièces jointes (nom unique dans le courriel, octets). Un message transféré devient un .eml."""
    out, used = [], set()
    for i, part in enumerate(msg.iter_attachments(), start=1):
        if part.get_content_type() == "message/rfc822":
            inner = part.get_content()
            payload = bytes(inner) if not isinstance(inner, bytes) else inner
            default = f"message_transfere_{i}.eml"
        else:
            payload = part.get_payload(decode=True) or b""
            ext = {"application/pdf": ".pdf", "text/plain": ".txt", "image/png": ".png"}.get(part.get_content_type(), ".bin")
            default = f"piece_jointe_{i}{ext}"
        name = (part.get_filename() or default).replace("\\", "/").rsplit("/", 1)[-1].strip() or default
        stem, dot, ext = name.rpartition(".") if "." in name else (name, "", "")
        n, unique = 2, name
        while unique in used:
            unique = f"{stem} ({n}){dot}{ext}" if dot else f"{name} ({n})"
            n += 1
        used.add(unique)
        out.append((unique, payload))
    return out


def parse_eml(raw: bytes):
    """Courriel : une unité pour le message ; les pièces jointes sont retournées à part."""
    msg = email.message_from_bytes(raw, policy=policy.default)
    body = msg.get_body(("plain", "html"))
    corps = ""
    if body is not None:
        corps = body.get_content()
        if body.get_content_subtype() == "html":
            corps = html_to_text(corps)
    corps = corps.strip()
    subject = str(msg["Subject"] or "")
    # En-tête Date non conforme RFC (client configuré en français : "mar., 8 sept. 2026 11:16") :
    # le parseur strict le vide, on relit la valeur brute.
    raw_date = next((str(v) for k, v in msg.raw_items() if k.lower() == "date"), "")
    date = email_date(raw_date) or find_date(raw_date)
    pjs = _attachments(msg)
    entete = f"Objet : {subject}\nDe : {msg['From'] or ''}\nÀ : {msg['To'] or ''}\nDate : {msg['Date'] or ''}"
    if pjs:
        entete += "\nPièces jointes : " + ", ".join(n for n, _ in pjs)
    unite = {"suffix": "msg", "unit": "courriel", "repere": f"courriel du {msg['Date'] or '?'}, de {msg['From'] or '?'}",
             "date_fait": date, "auteur": str(msg["From"] or "") or None, "texte": f"{entete}\n\n{corps}"}
    meta = {"doc_date": date, "title": subject, "message_id": norm_message_id(msg["Message-ID"]), "attachments": pjs}
    return meta, [unite]


# ───────────────────────── PDF ─────────────────────────

def parse_pdf(raw: bytes, transcrire=None):
    doc = fitz.open(stream=raw, filetype="pdf")
    unites, full, incomplet = [], [], False
    for n, page in enumerate(doc, start=1):
        txt, repere = page.get_text().strip(), f"p.{n}"
        if not txt:  # page scannée : pas de couche texte
            if transcrire:
                txt = transcrire(page.get_pixmap(dpi=150).tobytes("png"))
                repere = f"p.{n} (transcription IA, à vérifier)"
            else:
                txt = PLACEHOLDER
            incomplet |= txt == PLACEHOLDER
        full.append(txt)
        # « du 7 juillet au 31 octobre 2026 » : l'année du document complète les dates qui n'en ont pas
        annee = re.search(r"\b(20\d\d)\b", txt)
        unites.append({"suffix": f"p{n}", "unit": "page", "repere": repere,
                       "date_fait": find_date(txt, int(annee.group(1)) if annee else None),
                       "auteur": None, "texte": txt})
    lines = [l for l in "\n".join(full).splitlines() if l.strip()]
    title = next((l for l in lines if "fictif" not in l.lower() and "synthétique" not in l.lower()
                  and not l.lower().startswith("page")), "")
    return {"doc_date": next((u["date_fait"] for u in unites if u["date_fait"]), None), "title": title,
            "incomplet": incomplet}, unites


# ───────────────────────── Tableurs ─────────────────────────

def _fmt(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(timespec="minutes")
    if isinstance(v, dt.date):
        return v.isoformat()
    return str(v)


def parse_xlsx(raw: bytes):
    """Tableur : une unité par rangée, avec les en-têtes de colonnes pour le contexte.

    - l'en-tête est la première rangée avec au moins 2 cellules remplies (des titres peuvent la précéder) ;
    - une formule sans valeur calculée en cache est conservée telle quelle ;
    - un commentaire est conservé même sur une cellule vide.
    """
    wb_val = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    wb_for = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    unites = []
    for ws in wb_val:
        wf = wb_for[ws.title]
        rows = list(ws.iter_rows())

        def cell_text(c):
            f = wf[c.coordinate].value
            v = c.value
            if v is None and isinstance(f, str) and f.startswith("="):
                return f"{f} (formule, valeur non calculée)"
            if isinstance(f, str) and f.startswith("=") and v is not None:
                return f"{_fmt(v)} (formule {f})"
            return None if v in (None, "") else _fmt(v)

        def filled(row):
            return [c for c in row if cell_text(c) is not None]

        h_idx = next((i for i, r in enumerate(rows) if len(filled(r)) >= 2), None)
        headers = {}
        if h_idx is not None:
            headers = {c.column_letter: cell_text(c) or c.column_letter for c in rows[h_idx]}
        for i, row in enumerate(rows):
            parts = [f"{headers.get(c.column_letter, c.column_letter)} [{c.coordinate}]: {cell_text(c)}"
                     for c in filled(row)] if i != h_idx else []
            notes = [f"Commentaire {c.coordinate} ({wf[c.coordinate].comment.author or '?'}): {wf[c.coordinate].comment.text}"
                     for c in row if wf[c.coordinate].comment]
            if h_idx is not None and i < h_idx:
                parts = [cell_text(c) for c in filled(row)]   # rangée de titre
            if not parts and not notes:
                continue
            r = row[0].row
            unites.append({"suffix": f"{ws.title}!{r}", "unit": "rangee",
                           "repere": f"{ws.title}!A{r}:{row[-1].column_letter}{r}",
                           "date_fait": None, "auteur": None, "texte": "\n".join([" | ".join(parts)] + notes).strip()})
    return {"doc_date": None, "title": wb_val.sheetnames[0]}, unites


def parse_csv(text: str):
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(x.strip() for x in r)]
    if not rows:
        return {"doc_date": None, "title": ""}, []
    header, unites = rows[0], []
    for i, row in enumerate(rows[1:], start=2):
        texte = " | ".join(f"{h}: {v}" for h, v in zip(header, row))
        unites.append({"suffix": f"L{i}", "unit": "rangee", "repere": f"ligne {i}",
                       "date_fait": find_date(texte), "auteur": None, "texte": texte})
    return {"doc_date": None, "title": ", ".join(header)}, unites


def parse_docx(raw: bytes):
    import docx
    d = docx.Document(io.BytesIO(raw))
    paras = [p.text for p in d.paragraphs]
    for t in d.tables:
        paras += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
    return parse_lines("\n".join(paras), label="paragraphe")


# ───────────────────────── Images ─────────────────────────

VISION_PROMPT = (
    "Transcris fidèlement tout le texte visible de cette capture d'écran, ligne par ligne, "
    "en conservant les statuts et valeurs (ex. OK, TODO, codes d'erreur) et en signalant entre crochets "
    "ce que les couleurs indiquent (ex. [en rouge]). Termine par une ligne 'Description :' de 1 à 2 phrases "
    "sur ce que montre la capture. N'invente rien ; si un élément est illisible, écris [illisible]."
)


def parse_image(raw: bytes, transcrire):
    """Capture : transcription par un modèle de vision (à vérifier par une personne)."""
    texte = transcrire(raw) if transcrire else PLACEHOLDER
    return {"doc_date": find_date(texte), "title": "capture", "incomplet": texte == PLACEHOLDER}, [{
        "suffix": "capture", "unit": "capture", "repere": "capture (transcription IA, à vérifier)",
        "date_fait": find_date(texte), "auteur": None, "texte": texte}]


# ───────────────────────── Aiguillage ─────────────────────────

def parse(path_name: str, raw: bytes, transcrire=None):
    ext = path_name.rsplit(".", 1)[-1].lower() if "." in path_name.rsplit("/", 1)[-1] else ""
    if ext == "eml":
        return parse_eml(raw)
    if ext == "pdf":
        return parse_pdf(raw, transcrire)
    if ext in ("xlsx", "xlsm"):
        return parse_xlsx(raw)
    if ext == "docx":
        return parse_docx(raw)
    if ext in IMAGE_EXT:
        return parse_image(raw, transcrire)
    if ext == "csv":
        return parse_csv(decode_text(raw))
    if ext in TEXT_EXT or (ext == "" and not looks_binary(raw)):
        return parse_lines(decode_text(raw))
    raise NonSupporte(f"format .{ext or '?'} non pris en charge")
